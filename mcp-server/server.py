import logging
import os
from datetime import datetime
from functools import wraps
from typing import Optional

from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

from fastmcp import FastMCP
import database as db
from models import ExpenseInput, BudgetInput

MCP_TOKEN = os.getenv("BUDGET_MCP_TOKEN", "")

mcp = FastMCP(
    name="budget-db",
    instructions="Budget tracking database server. Provides tools to save and query expenses and budgets.",
)


def log_tool(fn):
    @wraps(fn)
    async def wrapper(*args, **kwargs):
        logger.info("[TOOL] %s called at %s | args=%s kwargs=%s", fn.__name__, datetime.now().isoformat(), args, kwargs)
        result = await fn(*args, **kwargs)
        logger.info("[TOOL] %s returned: %s", fn.__name__, str(result)[:200])
        return result
    return wrapper


@mcp.tool()
@log_tool
async def get_categories() -> list[dict]:
    """Return all expense categories currently configured."""
    return await db.fetch_categories()


@mcp.tool()
@log_tool
async def add_category(name: str, description: Optional[str] = None) -> dict:
    """Add a new expense category. Use this when an expense does not fit any existing category.
    Returns the new category or an error if it already exists."""
    name = name.strip().title()
    if not name:
        return {"error": "Category name cannot be empty"}
    return await db.insert_category(name, description)


@mcp.tool()
@log_tool
async def save_expense(
    amount: float,
    category: str,
    date: str,
    merchant: Optional[str] = None,
    description: Optional[str] = None,
    source: str = "text",
) -> dict:
    """Save a new expense. Returns the saved record with its id."""
    categories = await db.fetch_categories()
    valid_names = [c["name"] for c in categories]
    if category not in valid_names:
        return {
            "error": f"Unknown category '{category}'. Valid categories: {', '.join(valid_names)}. "
                     "If this is a new type of expense, call add_category first."
        }
    try:
        inp = ExpenseInput(
            amount=amount,
            category=category,
            merchant=merchant,
            description=description,
            date=date,
            source=source,
        )
    except Exception as e:
        return {"error": str(e)}
    return await db.insert_expense(
        inp.amount, inp.category, inp.merchant, inp.description, inp.date, inp.source
    )


@mcp.tool()
@log_tool
async def get_expenses(month: int, year: int, category: Optional[str] = None) -> list[dict]:
    """Return all expenses for the given month/year, optionally filtered by category."""
    if category:
        categories = await db.fetch_categories()
        valid_names = [c["name"] for c in categories]
        if category not in valid_names:
            return [{"error": f"Invalid category. Must be one of: {', '.join(valid_names)}"}]
    return await db.fetch_expenses(month, year, category)


@mcp.tool()
@log_tool
async def get_monthly_summary(month: int, year: int) -> dict:
    """Return total spend per category, overall total, and budget status for each category."""
    return await db.fetch_monthly_summary(month, year)


@mcp.tool()
@log_tool
async def delete_expense(expense_id: int) -> dict:
    """Delete an expense by id. Returns confirmation."""
    return await db.remove_expense(expense_id)


@mcp.tool()
@log_tool
async def set_budget(category: str, monthly_limit: float) -> dict:
    """Create or update a budget limit for a category. Returns the saved budget."""
    categories = await db.fetch_categories()
    valid_names = [c["name"] for c in categories]
    if category not in valid_names:
        return {"error": f"Unknown category '{category}'. Valid categories: {', '.join(valid_names)}"}
    try:
        inp = BudgetInput(category=category, monthly_limit=monthly_limit)
    except Exception as e:
        return {"error": str(e)}
    return await db.upsert_budget(inp.category, inp.monthly_limit)


@mcp.tool()
@log_tool
async def get_budgets() -> list[dict]:
    """Return all configured budgets."""
    return await db.fetch_budgets()


@mcp.tool()
@log_tool
async def get_budget_status(month: int, year: int) -> list[dict]:
    """Return each category's spend vs. its budget limit, with percentage used."""
    return await db.fetch_budget_status(month, year)


@mcp.tool()
@log_tool
async def export_to_excel(month: int, year: int) -> dict:
    """Generate a .xlsx file for the given month. Returns filepath and row count."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return {"error": "openpyxl is not installed"}

    expenses = await db.fetch_expenses(month, year)
    if not expenses:
        return {"error": f"No expenses found for {year}-{month:02d}"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    headers = ["ID", "Date", "Category", "Merchant", "Description", "Amount", "Source", "Created At"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="4472C4")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, exp in enumerate(expenses, 2):
        ws.cell(row=row_idx, column=1, value=exp.get("id"))
        ws.cell(row=row_idx, column=2, value=exp.get("date"))
        ws.cell(row=row_idx, column=3, value=exp.get("category"))
        ws.cell(row=row_idx, column=4, value=exp.get("merchant"))
        ws.cell(row=row_idx, column=5, value=exp.get("description"))
        ws.cell(row=row_idx, column=6, value=exp.get("amount"))
        ws.cell(row=row_idx, column=7, value=exp.get("source"))
        ws.cell(row=row_idx, column=8, value=exp.get("created_at"))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    exports_dir = os.path.join(os.path.dirname(os.getenv("DB_PATH", "./budget.db")), "exports")
    os.makedirs(exports_dir, exist_ok=True)
    filepath = os.path.join(exports_dir, f"budget_{year}_{month:02d}.xlsx")
    wb.save(filepath)
    logger.info("Exported %d rows to %s", len(expenses), filepath)
    return {"filepath": filepath, "rows": len(expenses)}


TOOL_MAP = {
    "get_categories": get_categories,
    "add_category": add_category,
    "save_expense": save_expense,
    "get_expenses": get_expenses,
    "get_monthly_summary": get_monthly_summary,
    "delete_expense": delete_expense,
    "set_budget": set_budget,
    "get_budgets": get_budgets,
    "get_budget_status": get_budget_status,
    "export_to_excel": export_to_excel,
}

if __name__ == "__main__":
    from starlette.applications import Starlette
    from starlette.requests import Request
    from starlette.responses import JSONResponse as StarletteJSON
    from starlette.routing import Mount, Route
    from contextlib import asynccontextmanager
    import uvicorn

    asgi_app = mcp.http_app(path="/mcp")

    async def health(request: Request):
        return StarletteJSON({"status": "ok"})

    async def api(request: Request):
        """Simple REST endpoint for dashboard use: POST /api {"tool": "...", "arguments": {...}}"""
        try:
            body = await request.json()
        except Exception:
            return StarletteJSON({"error": "Invalid JSON"}, status_code=400)
        tool_name = body.get("tool")
        arguments = body.get("arguments", {})
        fn = TOOL_MAP.get(tool_name)
        if fn is None:
            return StarletteJSON({"error": f"Unknown tool: {tool_name}"}, status_code=404)
        try:
            result = await fn(**arguments)
            return StarletteJSON(result)
        except Exception as e:
            logger.exception("Error calling tool %s", tool_name)
            return StarletteJSON({"error": str(e)}, status_code=500)

    @asynccontextmanager
    async def lifespan(app):
        await db.run_migrations()
        async with asgi_app.lifespan(app):
            logger.info("MCP Budget Server ready on http://localhost:8000")
            yield

    app = Starlette(
        lifespan=lifespan,
        routes=[
            Route("/health", health),
            Route("/api", api, methods=["POST"]),
            Mount("/", app=asgi_app),
        ],
    )

    uvicorn.run(app, host="0.0.0.0", port=8000)
